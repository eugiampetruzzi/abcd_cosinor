"""Build APA Table — Participant Characteristics."""
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os

OUT = os.path.dirname(os.path.abspath(__file__))
doc = Document()

style = doc.styles["Normal"]
style.font.name = "Arial"
style.font.size = Pt(12)
style.paragraph_format.space_after = Pt(0)
style.paragraph_format.space_before = Pt(0)
style.paragraph_format.line_spacing = 2.0

for section in doc.sections:
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)


def set_cell(cell, text, bold=False, italic=False, align="left", size=10):
    cell.text = ""
    p = cell.paragraphs[0]
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "right":
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = 1.15


def no_border(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for side in ["top", "left", "bottom", "right"]:
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "none")
        el.set(qn("w:sz"), "0")
        tcBorders.append(el)
    tcPr.append(tcBorders)


def border_bottom(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for side in ["top", "left", "right"]:
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "none")
        el.set(qn("w:sz"), "0")
        tcBorders.append(el)
    bot = OxmlElement("w:bottom")
    bot.set(qn("w:val"), "single")
    bot.set(qn("w:sz"), "4")
    bot.set(qn("w:color"), "000000")
    tcBorders.append(bot)
    tcPr.append(tcBorders)


def border_top(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for side in ["left", "right", "bottom"]:
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "none")
        el.set(qn("w:sz"), "0")
        tcBorders.append(el)
    top = OxmlElement("w:top")
    top.set(qn("w:val"), "single")
    top.set(qn("w:sz"), "4")
    top.set(qn("w:color"), "000000")
    tcBorders.append(top)
    tcPr.append(tcBorders)


def border_top_bottom(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for side in ["left", "right"]:
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "none")
        el.set(qn("w:sz"), "0")
        tcBorders.append(el)
    for side in ["top", "bottom"]:
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "000000")
        tcBorders.append(el)
    tcPr.append(tcBorders)


# --- Title ---
p = doc.add_paragraph()
run = p.add_run("Table X")
run.bold = True
run.font.name = "Arial"
run.font.size = Pt(12)

p2 = doc.add_paragraph()
run2 = p2.add_run("Participant Characteristics")
run2.italic = True
run2.font.name = "Arial"
run2.font.size = Pt(12)


# --- Data ---
# Col 0 = Variable, Col 1 = Full sample, Col 2 = Prediction sample
rows_data = [
    # (variable_text, full_val, pred_val, is_section_header, indent)
    ("Demographics", "", "", True, False),
    ("Sessions (participants)", "", "", False, False),
    ("  Wave 2", "7,230", "3,087", False, True),
    ("  Wave 4", "4,306", "", False, True),
    ("  Wave 6", "2,016", "", False, True),
    ("Age (years)", "", "", False, False),
    ("  Wave 2", "12.03 (0.66)", "12.00 (0.66)", False, True),
    ("  Wave 4", "14.18 (0.71)", "", False, True),
    ("  Wave 6", "16.08 (0.66)", "", False, True),
    ("Female", "4,061 (48.9)", "1,901 (61.6)", False, False),
    ("Race / ethnicity", "", "", False, False),
    ("  White", "4,760 (57.3)", "1,987 (64.4)", False, True),
    ("  Hispanic", "1,532 (18.5)", "442 (14.3)", False, True),
    ("  Black", "957 (11.5)", "262 (8.5)", False, True),
    ("  Asian", "152 (1.8)", "65 (2.1)", False, True),
    ("  Other", "870 (10.5)", "322 (10.4)", False, True),
    ("Household income", "", "", False, False),
    ("  < $50,000", "1,914 (23.1)", "537 (17.4)", False, True),
    ("  $50,000–$100,000", "2,317 (27.9)", "889 (28.8)", False, True),
    ("  > $100,000", "3,488 (42.0)", "1,487 (48.2)", False, True),
    ("Parental education", "", "", False, False),
    ("  Less than HS", "395 (4.8)", "108 (3.5)", False, True),
    ("  HS / GED", "696 (8.4)", "199 (6.4)", False, True),
    ("  Some college", "2,367 (28.5)", "756 (24.5)", False, True),
    ("  Bachelor’s degree", "2,556 (30.8)", "1,053 (34.1)", False, True),
    ("  Graduate / professional", "2,275 (27.4)", "968 (31.4)", False, True),

    ("Wave 2 cardiac rhythm", "", "", True, False),
    ("Mesor (bpm)", "82.37 (7.68)", "82.25 (7.51)", False, False),
    ("Amplitude (bpm)", "12.18 (3.01)", "12.32 (2.95)", False, False),
    ("Acrophase (clock hour)", "14.96 (1.26)", "14.86 (1.13)", False, False),

    ("Wave 2 clinical measures", "", "", True, False),
    ("CBCL Depression T-score", "54.27 (6.08)", "53.37 (4.73)", False, False),
    ("CBCL Anxiety T-score", "53.70 (6.05)", "53.05 (5.14)", False, False),
    ("CBCL ADHD T-score", "53.30 (5.42)", "53.03 (5.15)", False, False),
    ("CBCL Internalizing T-score", "47.81 (10.39)", "47.38 (9.44)", False, False),
    ("CBCL Externalizing T-score", "44.29 (9.66)", "43.95 (9.21)", False, False),
    ("BMI (kg/m²)", "20.51 (4.89)", "18.66 (3.48)", False, False),
    ("Systolic BP (mmHg)", "102.46 (10.75)", "101.47 (10.29)", False, False),
    ("Diastolic BP (mmHg)", "60.24 (8.63)", "59.62 (8.11)", False, False),
]

ncols = 3
nrows = len(rows_data) + 2  # +2 for header rows

table = doc.add_table(rows=nrows, cols=ncols)
table.alignment = WD_TABLE_ALIGNMENT.CENTER

# Header row 1  empty + sample labels
h0 = table.rows[0]
set_cell(h0.cells[0], "", align="left", size=10)
border_top(h0.cells[0])
set_cell(h0.cells[1], "Full sample", italic=True, align="center", size=10)
border_top(h0.cells[1])
border_bottom(h0.cells[1])
set_cell(h0.cells[2], "Prediction sample", italic=True, align="center", size=10)
border_top(h0.cells[2])
border_bottom(h0.cells[2])

# Header row 2  N labels
h1 = table.rows[1]
set_cell(h1.cells[0], "", align="left", size=10)
border_bottom(h1.cells[0])
set_cell(h1.cells[1], "(N = 8,301)", italic=True, align="center", size=10)
border_bottom(h1.cells[1])
set_cell(h1.cells[2], "(N = 3,087)", italic=True, align="center", size=10)
border_bottom(h1.cells[2])

# Data rows
for i, (var, full, pred, is_header, indent) in enumerate(rows_data):
    row = table.rows[i + 2]
    set_cell(row.cells[0], var, bold=is_header, italic=is_header,
             align="left", size=10)
    set_cell(row.cells[1], full, align="center", size=10)
    set_cell(row.cells[2], pred, align="center", size=10)
    for j in range(ncols):
        no_border(row.cells[j])

# Bottom border on last row
for j in range(ncols):
    border_bottom(table.rows[-1].cells[j])

# Note
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(4)
run_label = p.add_run("Note. ")
run_label.italic = True
run_label.font.name = "Arial"
run_label.font.size = Pt(10)
run_text = p.add_run(
    "Continuous variables are M (SD); categorical variables are n (%). "
    "The full sample includes all participants contributing valid Fitbit data "
    "at one or more waves (Aim 1). The prediction sample is the union of "
    "depression, obesity, and hypertension analytic frames (Aim 2): participants "
    "who were below clinical threshold at baseline and had follow-up data. "
    "Wave 2 rhythm parameters are from the cosinor model. "
    "CBCL = Child Behavior Checklist; ADHD = Attention-Deficit/Hyperactivity "
    "Disorder; BMI = body mass index; BP = blood pressure; "
    "HS = high school; GED = General Educational Development."
)
run_text.font.name = "Arial"
run_text.font.size = Pt(10)

out_path = os.path.join(OUT, "Table_Participants.docx")
doc.save(out_path)
print(f"Saved to {out_path}")
